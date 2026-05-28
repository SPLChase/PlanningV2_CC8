from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from planning_v2.template_specs import template_columns


class TemplateSpecTests(unittest.TestCase):
    def test_individual_template_workbook_overrides_raw_schema(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            samples = Path(tmp)
            raw = Workbook()
            ws = raw.active
            ws.title = "WarehouseStockOnHand_Template"
            ws.append(["FieldName", "FieldType", "Examples", "Notes"])
            ws.append(["partCode", "String", "", ""])
            raw.save(samples / "Templates raw.xlsx")

            individual = Workbook()
            sheet = individual.active
            sheet.append(["partNumber", "warehouseCode"])
            sheet.append(["A1", "WH1"])
            individual.save(samples / "Warehouse Stock on Hand.xlsx")

            columns = template_columns(samples)

        self.assertEqual(columns["WarehouseStockOnHand"], ["partNumber", "warehouseCode"])

    def test_addresses_template_repairs_bad_sample_header(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            samples = Path(tmp)
            raw = Workbook()
            raw_ws = raw.active
            raw_ws.title = "Addresses_Template"
            raw_ws.append(["FieldName", "FieldType", "Examples", "Notes"])
            raw_ws.append(["externalAddressId", "String", "", ""])
            raw_ws.append(["addressLine1", "String", "", ""])
            raw.save(samples / "Templates raw.xlsx")

            individual = Workbook()
            sheet = individual.active
            sheet.append(["externalAddressId", "6,0", "addressLine2"])
            individual.save(samples / "Addresses.xlsx")

            columns = template_columns(samples)

        self.assertEqual(columns["Addresses"], ["externalAddressId", "addressLine1", "addressLine2"])


if __name__ == "__main__":
    unittest.main()
