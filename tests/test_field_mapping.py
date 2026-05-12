from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from planning_v2.field_mapping import (
    STATUS_CONFIRMED,
    STATUS_INVESTIGATE_SAP,
    clean_text,
    classify_field,
    parse_target_fields,
)


class FieldMappingTests(unittest.TestCase):
    def test_clean_text_removes_bom_and_nullish_values(self) -> None:
        self.assertEqual(clean_text("\ufeffpart_id "), "part_id")
        self.assertEqual(clean_text(None), "")
        self.assertEqual(clean_text("nan"), "")

    def test_parse_target_fields_reads_required_reference_columns(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            ref_dir = Path(tmp)
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(
                [
                    "Field name",
                    "Key column",
                    "Example data 1",
                    "Example data 2",
                    "Example data 3",
                    "Used by SPL",
                    "Data Type Info",
                    "Field Explanation",
                    "Context Area",
                    "Altron file",
                    "CC8 relevant?",
                    "CC8 can supply?",
                    "Priority",
                    "CC8Comment",
                    "Codex",
                ]
            )
            worksheet.append(["\ufeffpart_id", "** KEY **", "", "", "", "YES", "String", "Part", "SPL Masters", "AMS Masters (SPL generated - from distribution portal)", "Yes", "Yes", 1, "Use masters", ""])
            workbook.save(ref_dir / "SPL Planning Data Fields.xlsx")

            fields = parse_target_fields(ref_dir)

        self.assertEqual(len(fields), 1)
        self.assertEqual(fields[0].field_name, "part_id")
        self.assertEqual(fields[0].priority, "1")
        self.assertEqual(fields[0].cc8_comment, "Use masters")

    def test_classify_known_minstock_master_field(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            ref_dir = Path(tmp)
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(
                [
                    "Field name",
                    "Key column",
                    "Example data 1",
                    "Example data 2",
                    "Example data 3",
                    "Used by SPL",
                    "Data Type Info",
                    "Field Explanation",
                    "Context Area",
                    "Altron file",
                    "CC8 relevant?",
                    "CC8 can supply?",
                    "Priority",
                    "CC8Comment",
                    "Codex",
                ]
            )
            worksheet.append(["SPL MASTER", "", "", "", "", "YES", "String", "SPL master", "SPL Masters", "AMS Masters (SPL generated - from distribution portal)", "Yes", "Yes", 1, "", ""])
            workbook.save(ref_dir / "SPL Planning Data Fields.xlsx")
            field = parse_target_fields(ref_dir)[0]

        decision = classify_field(field)
        self.assertEqual(decision.source_status, STATUS_CONFIRMED)
        self.assertIn("masters.csv", decision.source_object)

    def test_classify_power_automate_hint_as_sap_first_investigation(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            ref_dir = Path(tmp)
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(
                [
                    "Field name",
                    "Key column",
                    "Example data 1",
                    "Example data 2",
                    "Example data 3",
                    "Used by SPL",
                    "Data Type Info",
                    "Field Explanation",
                    "Context Area",
                    "Altron file",
                    "CC8 relevant?",
                    "CC8 can supply?",
                    "Priority",
                    "CC8Comment",
                    "Codex",
                ]
            )
            worksheet.append(["order_id", "", "", "", "", "YES", "String", "Work order", "Work order information", "order_line", "Yes", "Via Chase Power Automate", 1, "Maybe from tickets", ""])
            workbook.save(ref_dir / "SPL Planning Data Fields.xlsx")
            field = parse_target_fields(ref_dir)[0]

        decision = classify_field(field)
        self.assertEqual(decision.source_status, STATUS_INVESTIGATE_SAP)


if __name__ == "__main__":
    unittest.main()
