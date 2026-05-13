"""Planning V2 template schema helpers."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

TEMPLATES_WORKBOOK = "Templates raw.xlsx"

OUTPUT_TO_TEMPLATE_FILE = {
    "Nodes": "Nodes.xlsx",
    "Customers": "Customers.xlsx",
    "Vendors": "Vendors.xlsx",
    "Parts": "Parts.xlsx",
    "Addresses": "Addresses.xlsx",
    "Warehouses": "Warehouses.xlsx",
    "WarehouseStockOnHand": "Warehouse Stock on Hand.xlsx",
    "Employees": "Employees.xlsx",
    "PartsUsage": "PartsUsage.xlsx",
    "ServiceOrder": "ServiceOrder.xlsx",
    "RepairOrder": "RepairOrders.xlsx",
    "PartCost": "PartCosts.xlsx",
    "Models": "Models.xlsx",
    "InventoryTransfers": "Inventory Transfers.xlsx",
    "PurchaseOrders": "Purchase Orders.xlsx",
    "PartTypes": "partTypes.xlsx",
    "Yields": "Yields.xlsx",
    "ActionGroups": "ActionGroups.xlsx",
    "WarehouseExclusions": "Warehouse Distribution Exclusions.xlsx",
}

_TEMPLATE_FIELD_CACHE: dict[str, dict[str, list[TemplateField]]] = {}

SHEET_TO_OUTPUT = {
    "Nodes_Template": "Nodes",
    "Customers_Template": "Customers",
    "Vendors_Template": "Vendors",
    "Parts_Template": "Parts",
    "Addresses_Template": "Addresses",
    "Warehouses_Template": "Warehouses",
    "WarehouseStockOnHand_Template": "WarehouseStockOnHand",
    "Employees_Template": "Employees",
    "PartsUsage_Template": "PartsUsage",
    "ServiceOrder_Template": "ServiceOrder",
    "RepairOrder_Template": "RepairOrder",
    "PartCost_Template": "PartCost",
    "Models_Template": "Models",
    "InventoryTransfers_Template": "InventoryTransfers",
    "PurchaseOrders_Template": "PurchaseOrders",
    "PartTypes_Template": "PartTypes",
    "Yields_Template": "Yields",
    "ActionGroups_Template": "ActionGroups",
    "WarehouseExclusions_Template": "WarehouseExclusions",
}


@dataclass(frozen=True)
class TemplateField:
    output_object: str
    field_name: str
    field_type: str
    examples: str
    notes: str


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\ufeff", "").strip()
    if text.lower() in {"none", "nan", "null", "<na>"}:
        return ""
    return text


def normalize_object_name(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")


def load_template_fields(samples_dir: Path) -> dict[str, list[TemplateField]]:
    cache_key = str(samples_dir.resolve())
    if cache_key in _TEMPLATE_FIELD_CACHE:
        return _TEMPLATE_FIELD_CACHE[cache_key]
    raw_templates = _load_raw_template_fields(samples_dir)
    templates: dict[str, list[TemplateField]] = {}
    for output_object, file_name in OUTPUT_TO_TEMPLATE_FILE.items():
        path = samples_dir / file_name
        if path.exists():
            headers = _load_individual_template_headers(path)
            raw_by_name = {field.field_name: field for field in raw_templates.get(output_object, [])}
            templates[output_object] = [
                TemplateField(
                    output_object=output_object,
                    field_name=header,
                    field_type=raw_by_name.get(header, TemplateField(output_object, header, "", "", "")).field_type,
                    examples=raw_by_name.get(header, TemplateField(output_object, header, "", "", "")).examples,
                    notes=f"Canonical header from {file_name}",
                )
                for header in headers
            ]
        else:
            templates[output_object] = raw_templates.get(output_object, [])
    if not templates:
        raise FileNotFoundError(f"Missing Planning V2 templates under: {samples_dir}")
    _TEMPLATE_FIELD_CACHE[cache_key] = templates
    return templates


def _load_individual_template_headers(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    workbook.close()
    return [clean_text(value) for value in row if clean_text(value)]


def _load_raw_template_fields(samples_dir: Path) -> dict[str, list[TemplateField]]:
    path = samples_dir / TEMPLATES_WORKBOOK
    if not path.exists():
        return {}

    workbook = load_workbook(path, read_only=True, data_only=True)
    templates: dict[str, list[TemplateField]] = {}
    for sheet_name, output_object in SHEET_TO_OUTPUT.items():
        if sheet_name not in workbook.sheetnames:
            templates[output_object] = []
            continue
        worksheet = workbook[sheet_name]
        fields: list[TemplateField] = []
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            field_name = clean_text(row[0] if len(row) > 0 else "")
            if not field_name:
                continue
            fields.append(
                TemplateField(
                    output_object=output_object,
                    field_name=field_name,
                    field_type=clean_text(row[1] if len(row) > 1 else ""),
                    examples=clean_text(row[2] if len(row) > 2 else ""),
                    notes=clean_text(row[3] if len(row) > 3 else ""),
                )
            )
        templates[output_object] = fields
    workbook.close()
    return templates


def template_columns(samples_dir: Path) -> dict[str, list[str]]:
    return {
        output_object: [field.field_name for field in fields]
        for output_object, fields in load_template_fields(samples_dir).items()
    }
