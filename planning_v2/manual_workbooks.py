"""Small human-fill workbook builders for CoCre8 onboarding."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


WAREHOUSE_FILL_COLUMNS = [
    "warehouseId",
    "warehouseDescription",
    "addressId",
    "returnWarehouseId",
    "supplyWarehouseId",
    "warehouseTypeId",
    "isReplenishable",
    "isBranchStockable",
    "isRemote",
    "warehouseStatusId",
    "isObsolete",
    "activityEvidence",
    "Notes",
]


def read_existing_warehouse_fill(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame(columns=WAREHOUSE_FILL_COLUMNS)
    try:
        df = pd.read_excel(path, sheet_name="Warehouses To Fill", dtype=str).fillna("")
    except Exception:
        return pd.DataFrame(columns=WAREHOUSE_FILL_COLUMNS)
    for column in WAREHOUSE_FILL_COLUMNS:
        if column not in df.columns:
            df[column] = ""
    return df[WAREHOUSE_FILL_COLUMNS]


def build_warehouse_fill_workbook(
    warehouses_csv: Path = Path("data/output/onboarding_csvs/Warehouses.csv"),
    existing_fill_path: Path = Path("docs/manual-fill/Warehouses_Missing_Fields_To_Fill.xlsx"),
    out_path: Path = Path("docs/manual-fill/Warehouses_Missing_Fields_To_Fill_simple.xlsx"),
) -> Path:
    warehouses = pd.read_csv(warehouses_csv, dtype=str).fillna("")
    existing = read_existing_warehouse_fill(existing_fill_path)
    existing_by_id = existing.set_index("warehouseId").to_dict("index") if not existing.empty else {}

    rows: list[dict[str, str]] = []
    for _, source in warehouses.iterrows():
        warehouse_id = str(source.get("warehouseId", "")).strip()
        if not warehouse_id:
            continue
        previous = existing_by_id.get(warehouse_id, {})
        row = {column: str(previous.get(column, "") or "") for column in WAREHOUSE_FILL_COLUMNS}
        row["warehouseId"] = warehouse_id
        row["warehouseDescription"] = str(source.get("warehouseDescription", "")).strip()
        row["addressId"] = warehouse_id
        row["isObsolete"] = str(previous.get("isObsolete", "") or "")
        row["activityEvidence"] = str(previous.get("activityEvidence", "") or "")
        rows.append(row)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Warehouses To Fill"
    worksheet.append(WAREHOUSE_FILL_COLUMNS)
    for row in rows:
        worksheet.append([row.get(column, "") for column in WAREHOUSE_FILL_COLUMNS])

    _format_fill_sheet(worksheet, WAREHOUSE_FILL_COLUMNS)
    _add_field_notes(workbook)
    _add_warehouse_ids(workbook, rows)
    workbook.save(out_path)
    check = load_workbook(out_path, read_only=True, data_only=True)
    check.close()
    return out_path


def _format_fill_sheet(worksheet, columns: list[str]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    manual_fill = PatternFill("solid", fgColor="FFF2CC")
    locked_fill = PatternFill("solid", fgColor="D9EAD3")
    derived_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            header = worksheet.cell(row=1, column=cell.column).value
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if header in {"warehouseId", "warehouseDescription"}:
                cell.fill = locked_fill
            elif header == "addressId":
                cell.fill = derived_fill
            else:
                cell.fill = manual_fill
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    for index, column in enumerate(columns, start=1):
        width = max(len(column), *(len(str(worksheet.cell(row, index).value or "")) for row in range(2, worksheet.max_row + 1)))
        worksheet.column_dimensions[get_column_letter(index)].width = min(width + 2, 45)
    for column_name in ["isReplenishable", "isBranchStockable", "isRemote"]:
        index = columns.index(column_name) + 1
        validation = DataValidation(type="list", formula1='"TRUE,FALSE,N/A"', allow_blank=True)
        worksheet.add_data_validation(validation)
        validation.add(f"{get_column_letter(index)}2:{get_column_letter(index)}{worksheet.max_row}")
    if "isObsolete" in columns:
        index = columns.index("isObsolete") + 1
        validation = DataValidation(type="list", formula1='"Y,N,N/A"', allow_blank=True)
        worksheet.add_data_validation(validation)
        validation.add(f"{get_column_letter(index)}2:{get_column_letter(index)}{worksheet.max_row}")


def _add_field_notes(workbook: Workbook) -> None:
    worksheet = workbook.create_sheet("Field Notes")
    notes = [
        ("addressId", "Derived from warehouseId / SAP warehouse code per user mapping. Example: FUJ CT."),
        ("returnWarehouseId", "Preferred warehouse for returned/faulty parts. Use a warehouseId from the WarehouseIds sheet, or N/A."),
        ("supplyWarehouseId", "Preferred replenishment/source warehouse. Blank means virtual/quarantine/faulty-return warehouse and is excluded from MVP planning outputs."),
        ("warehouseTypeId", "Planning warehouse type. Examples in template are BOOT, DEP, REPAIR. Use a short controlled value we can standardise later."),
        ("isReplenishable", "TRUE if this warehouse should be replenished by the planning engine; FALSE/N/A otherwise."),
        ("isBranchStockable", "TRUE if this branch/location can hold branch stock; FALSE/N/A otherwise."),
        ("isRemote", "TRUE if this is an outlying/remote location; FALSE/N/A otherwise."),
        ("warehouseStatusId", "Derived in upload as active flag: Y when isObsolete is N, N when isObsolete is Y. Leave this blank in the workbook unless manually overriding later."),
        ("isObsolete", "Use Y when the warehouse/site is not being used, empty, or should not participate in planning. Use N for active planning warehouses."),
        ("activityEvidence", "Optional evidence/comment for recent stock movement or empty status. This is for review and is not a template field."),
    ]
    worksheet.append(["Field", "What to Fill"])
    for note in notes:
        worksheet.append(note)
    _format_notes_sheet(worksheet, 110)


def _add_warehouse_ids(workbook: Workbook, rows: list[dict[str, str]]) -> None:
    worksheet = workbook.create_sheet("WarehouseIds")
    worksheet.append(["warehouseId", "warehouseDescription"])
    seen: set[str] = set()
    for row in rows:
        warehouse_id = row.get("warehouseId", "")
        if not warehouse_id or warehouse_id in seen:
            continue
        seen.add(warehouse_id)
        worksheet.append([warehouse_id, row.get("warehouseDescription", "")])
    _format_notes_sheet(worksheet, 55)


def _format_notes_sheet(worksheet, second_col_width: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = second_col_width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
