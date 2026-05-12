"""Build the CoCre8 Planning V2 field map workbook."""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from planning_v2.config import PlanningConfig, get_config
from planning_v2.schemas import CONFIRMED_OUTPUT_OBJECTS, FIELD_MAP_WORKBOOK, PENDING_OUTPUT_OBJECTS
from planning_v2.template_specs import load_template_fields, template_columns
from planning_v2.onboarding_csvs import POPULATED_TEMPLATE_FIELDS


TARGET_FIELDS_FILE = "SPL Planning Data Fields edited.xlsx"
FALLBACK_TARGET_FIELDS_FILE = "SPL Planning Data Fields.xlsx"

STATUS_CONFIRMED = "Confirmed"
STATUS_REVIEW = "Review required"
STATUS_INVESTIGATE_SAP = "Investigate SAP first"
STATUS_INVESTIGATE_EXTERNAL = "Investigate source"
STATUS_NOT_AVAILABLE = "Not available in checked source"
STATUS_OUT_OF_SCOPE = "Out of v1 scope"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
CONFIRMED_FILL = PatternFill("solid", fgColor="D9EAD3")
REVIEW_FILL = PatternFill("solid", fgColor="FCE4D6")
UNKNOWN_FILL = PatternFill("solid", fgColor="F4CCCC")
STATIC_FILL = PatternFill("solid", fgColor="E7E6E6")
TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)

WAREHOUSE_TEMPLATE_NOTES = {
    "addressId": "Expected site/address id linking to the Addresses template. Data dictionary field: warehouse.address_id.",
    "nodeId": "Expected associated planning hierarchy node id. Data fields list this as warehouse.node_id / associated node.",
    "returnWarehouseId": "Expected preferred warehouse for return/faulty stock going to repair or scrap. Data fields: return_warehouse.",
    "supplyWarehouseId": "Expected preferred replenishment/source warehouse. Data fields: supply_warehouse.",
    "warehouseTypeId": "Expected site type enum; sample template uses BOOT, DEP, and REPAIR. Data fields mark warehouse_type_id as critical.",
    "isRepairWarehouse": "Expected true when this is an internal repair location. Data dictionary field: is_repair_whse.",
    "isReplenishable": "Expected true when the site can be restocked. Data dictionary field: is_replenishable.",
    "isBootStockable": "Expected true when technician/boot location can accept stock. Data fields: cst_boot_stockable.",
    "isBranchStockable": "Expected true when branch/regional location can accept stock. Data fields: cst_branch_stockable.",
    "isRemote": "Expected true when this is an outlying/regional location. Data fields: cst_remote.",
    "warehouseStatusId": "Expected current/inactive status if required by the importer. Template says confirm; data fields mention cst_active and is_obsolete.",
}

CUSTOMER_TEMPLATE_NOTES = {
    "customerId": "Left blank. SAP BusinessPartners can supply customer codes, but the correct Planning V2 customer scope is not confirmed.",
    "customerGroupId": "Left blank. User has no confirmed source yet for customer grouping.",
    "Description": "Left blank. SAP BusinessPartners can supply names, but the correct Planning V2 customer scope is not confirmed.",
    "assignAnySkill": "Left blank. User has no confirmed source yet for customer skill assignment rules.",
    "isActive": "Left blank. User has no confirmed source yet for customer active/inactive status.",
    "dseSlaCost": "Left blank. User has no confirmed source yet for SLA cost.",
    "dseSlaRevenue": "Left blank. User has no confirmed source yet for SLA revenue.",
    "stdResponseTime": "Left blank. User has no confirmed source yet for standard response time.",
    "stdRepairTime": "Left blank. User has no confirmed source yet for standard repair time.",
}

STOCK_ON_HAND_TEMPLATE_NOTES = {
    "inventoryType": "Left blank. SAP OITW live stock metrics do not identify Planning V2 good/bad inventory type; do not infer from warehouse names.",
    "quantityOutbound": "Left blank. SAP OITW has IsCommited, which is already mapped to quantityAllocated; no separate outbound quantity was found in the live stock metrics.",
    "uniqueId": "Left blank. Template expects an integer technical id, but no confirmed source/key rule has been provided.",
}

TEMPLATE_FIELDS_NOT_CC8_RELEVANT = {
    ("ActionGroups", "actionGroupId"): "actgr.actgr_id is marked CC8 relevant = No.",
    ("ActionGroups", "nodeId"): "actgr.node_id is marked CC8 relevant = No.",
    ("ActionGroups", "actionGroupDescription"): "actgr.descr is marked CC8 relevant = No.",
    ("ActionGroups", "assignAnySkill"): "Technician action groups are marked CC8 relevant = No.",
    ("ActionGroups", "isUsed"): "actgr.is_used is marked CC8 relevant = No.",
    ("ActionGroups", "isObsolete"): "actgr.is_obsolete is marked CC8 relevant = No.",
    ("Employees", "actionGroupId"): "person.actgr_id is marked CC8 relevant = No.",
    ("InventoryTransfers", "loClass"): "imo.LO_Class is marked CC8 relevant = No.",
    ("PartCost", "averageRepairCost"): "bpart_cost.avg_repair_cost is marked CC8 relevant = No.",
    ("Parts", "isService"): "bpart.is_service is marked CC8 relevant = No.",
    ("Parts", "isTool"): "bpart.is_tool is marked CC8 relevant = No.",
    ("Parts", "isSmallPart"): "bpart.cst_small_part is marked CC8 relevant = No.",
    ("PartsUsage", "requestId"): "parts_usage.Request_ID is marked CC8 relevant = No.",
    ("PurchaseOrders", "customerId"): "No corresponding purchase_orders field is marked needed for CoCre8 in the edited reference workbook.",
    ("PurchaseOrders", "requestTicketDateTime"): "No corresponding purchase_orders field is marked needed for CoCre8 in the edited reference workbook.",
    ("PurchaseOrders", "isResolved"): "No corresponding purchase_orders field is marked needed for CoCre8 in the edited reference workbook.",
    ("Warehouses", "nodeId"): "warehouse.node_id is marked CC8 relevant = No.",
    ("Warehouses", "isRepairWarehouse"): "warehouse.is_repair_whse is marked CC8 relevant = No.",
    ("Warehouses", "isBootStockable"): "warehouse.cst_boot_stockable is marked CC8 relevant = No.",
}


@dataclass(frozen=True)
class TargetField:
    row_number: int
    field_name: str
    object_name: str
    context_area: str
    key_column: str
    used_by_spl: str
    data_type: str
    explanation: str
    cc8_relevant: str
    cc8_can_supply: str
    priority: str
    cc8_comment: str
    codex: str


@dataclass(frozen=True)
class FieldDecision:
    target: TargetField
    source_status: str
    source_system: str
    source_object: str
    source_field: str
    transform: str
    confidence: str
    notes: str
    codex_note: str = ""


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\ufeff", "").strip()
    if text.lower() in {"none", "nan", "null", "<na>"}:
        return ""
    return text


def normalize_name(value: object) -> str:
    text = clean_text(value).lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def parse_target_fields(reference_dir: Path) -> list[TargetField]:
    path = reference_dir / TARGET_FIELDS_FILE
    if not path.exists():
        path = reference_dir / FALLBACK_TARGET_FIELDS_FILE
    if not path.exists():
        raise FileNotFoundError(f"Missing target field workbook: {reference_dir / TARGET_FIELDS_FILE}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["SPL Planning Data Fields"] if "SPL Planning Data Fields" in workbook.sheetnames else workbook.active
    headers = [clean_text(cell) for cell in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    index = {header: i for i, header in enumerate(headers) if header}
    required = [
        "Field name",
        "Key column",
        "Used by SPL",
        "Data Type Info",
        "Field Explanation",
        "Context Area",
        "Altron file",
        "CC8 relevant?",
        "CC8 can supply?",
    ]
    missing = [name for name in required if name not in index]
    if missing:
        raise ValueError(f"Missing columns in {path}: {', '.join(missing)}")

    fields: list[TargetField] = []
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        field_name = clean_text(row[index["Field name"]])
        if not field_name:
            continue
        fields.append(
            TargetField(
                row_number=row_number,
                field_name=field_name,
                object_name=clean_text(row[index["Altron file"]]),
                context_area=clean_text(row[index["Context Area"]]),
                key_column=clean_text(row[index["Key column"]]),
                used_by_spl=clean_text(row[index["Used by SPL"]]),
                data_type=clean_text(row[index["Data Type Info"]]),
                explanation=clean_text(row[index["Field Explanation"]]),
                cc8_relevant=clean_text(row[index["CC8 relevant?"]]),
                cc8_can_supply=clean_text(row[index["CC8 can supply?"]]),
                priority=clean_text(row[index["Priority"]]) if "Priority" in index else "",
                cc8_comment=clean_text(row[index["CC8Comment"]]) if "CC8Comment" in index else "",
                codex=clean_text(row[index["Codex"]]) if "Codex" in index else "",
            )
        )
    workbook.close()
    return fields


def is_v1_target(field: TargetField) -> bool:
    relevance = field.cc8_relevant.lower()
    priority = priority_rank(field)
    if priority == 0:
        return False
    if priority in {1, 2, 3}:
        return not relevance.startswith("no")
    if relevance.startswith("no"):
        return False
    if relevance.startswith("yes") or relevance.startswith("maybe"):
        return True
    if relevance.startswith("not sure"):
        return True
    if field.context_area in {"Stock detail", "warehouse details"}:
        return field.used_by_spl.upper() == "YES"
    if field.context_area in {"purchase_orders", "repair_orders"}:
        return field.used_by_spl.upper() == "YES"
    return False


def priority_rank(field: TargetField) -> int | None:
    text = clean_text(field.priority)
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def classify_field(field: TargetField) -> FieldDecision:
    object_key = normalize_name(field.object_name)
    field_key = normalize_name(field.field_name)
    supplied = field.cc8_can_supply.lower()
    context = field.context_area
    priority = priority_rank(field)

    if priority == 0 or field.cc8_relevant.lower().startswith("no"):
        return FieldDecision(field, STATUS_OUT_OF_SCOPE, "", "", "", "", "", "Not marked as needed for CoCre8 onboarding.", "Priority 0 / CC8 not relevant; excluded from MVP blockers.")

    confirmed_sources = {
        "ams_masters_spl_generated_from_distribution_portal": ("Local file", "masters.csv", "SPL Master / Items linked"),
        "bpart_alt": ("Local file", "masters.csv", "Items linked"),
        "bpart_cost": ("SPI file", "SPI_DATA.csv", field.field_name),
        "whse_bpart_qty": ("SAP Service Layer", "SQLQueries on OITW/OITM", field.field_name),
        "warehouse": ("SAP Service Layer", "Warehouses endpoint", field.field_name),
        "cust": ("Local dimension", "customers.csv / warehouse_dimension.csv", field.field_name),
    }
    if object_key in confirmed_sources:
        source_system, source_object, source_field = confirmed_sources[object_key]
        return FieldDecision(
            target=field,
            source_status=STATUS_CONFIRMED,
            source_system=source_system,
            source_object=source_object,
            source_field=source_field,
            transform=_default_transform(field),
            confidence="High" if object_key != "bpart_cost" else "Medium",
            notes="Confirmed from MinStock3/Exco local sources; review business meaning where noted.",
            codex_note="Confirmed source. Use masters.csv as the SPL Master bridge; keep actual part numbers for traceability.",
        )

    if object_key == "bpart":
        if field_key in {"abc_id", "is_exclude_from_replishment", "cst_branch_stockable", "cst_boot_stock"}:
            return _review_decision(field, "Local/SPI-derived item attributes need business confirmation.")
        if field_key in {"part_id", "pclass2_id", "pclass3_id"}:
            return FieldDecision(field, STATUS_CONFIRMED, "SAP/SPI/local", "OITM / SPI_DATA.csv", field.field_name, _default_transform(field), "Medium", "Part identity and class-like attributes are present, but exact V2 semantic mapping needs review.", "Available, but use SPL Master as planning grain and actual SAP part as evidence.")
        return _sap_investigation_decision(field, "SAP item master field candidate; exact source column not yet proven.")

    if object_key in {"goods", "ib_warehouse", "model"}:
        return _sap_investigation_decision(field, "Installed base/model/BOM source not proven for CoCre8.")

    if object_key == "imo":
        return _sap_investigation_decision(field, "Inventory transfer or stock transfer request source must be verified in SAP.")

    if object_key == "order_line":
        return _sap_investigation_decision(field, "Marked Power Automate in reference, but plan requires SAP-first investigation.")

    if object_key == "parts_usage":
        if supplied.startswith("yes"):
            return FieldDecision(field, STATUS_CONFIRMED, "Manual SAP stock audit report", "Stock Audit Report", field.field_name, _default_transform(field), "Medium", "Usage is allowed from the manually generated SAP Stock Audit Report; work-order linkage may need SAP verification.")
        return _sap_investigation_decision(field, "Usage detail exists, but serial/date granularity needs investigation.")

    if object_key in {"person", "node", "demand_trans_log"}:
        return _sap_investigation_decision(field, "Reference marks this as possible but source is not proven.")

    if object_key in {"multiple", "critical_parts", "exclusions_import_template_spl_generated_by_distribution_portal"}:
        return FieldDecision(field, STATUS_INVESTIGATE_EXTERNAL, "SPL distribution portal/manual rule", field.object_name, field.field_name, "Import as maintained rule table once provided.", "Low", "Not available from MinStock3/Exco evidence.", "Needs non-SAP owner/source confirmation before population.")

    if object_key in {"purchase_orders", "repair_orders", "service_call_env"}:
        return _sap_investigation_decision(field, "Important V2 vision object; not confirmed in current CC8 local sources.")

    if supplied.startswith("yes"):
        return FieldDecision(field, STATUS_REVIEW, "Reference workbook", field.object_name, field.field_name, _default_transform(field), "Medium", "Reference says CC8 can supply, but no local proof was identified.", "CC8 note says available, but implementation needs proof/source column before loading.")
    if "power automate" in supplied:
        return _sap_investigation_decision(field, "External hint exists, but SAP-first investigation is required.")
    if is_v1_target(field):
        return _sap_investigation_decision(field, "V1 target field has no confirmed source.")
    return FieldDecision(field, STATUS_OUT_OF_SCOPE, "", "", "", "", "", "Not marked CC8 relevant for v1.", "Excluded from current CoCre8 onboarding scope.")


def _review_decision(field: TargetField, notes: str) -> FieldDecision:
    return FieldDecision(field, STATUS_REVIEW, "Local/SPI/SPL rule", field.object_name, field.field_name, _default_transform(field), "Medium", notes, "Review business meaning before loading; do not infer from field name alone.")


def _sap_investigation_decision(field: TargetField, notes: str) -> FieldDecision:
    return FieldDecision(field, STATUS_INVESTIGATE_SAP, "SAP Service Layer", field.object_name, field.field_name, "Do not populate until SAP source is proven.", "Low", notes, "Investigate SAP first; leave blank until exact endpoint/table and semantics are proven.")


def _default_transform(field: TargetField) -> str:
    key = normalize_name(field.field_name)
    if key in {"primary_part"}:
        return "Use first linked item from SPL master mapping; review required."
    if "cost" in key or "price" in key:
        return "Coerce numeric currency value; preserve original currency where available."
    if key in {"qty", "quantity", "qty_used", "allocated_qty", "in_bound_qty", "out_bound_qty"}:
        return "Coerce to numeric quantity."
    if "date" in key or "time" in key or key.endswith("_at"):
        return "Parse to ISO datetime/date where source supports it."
    return "Direct map after trimming text."


def build_decisions(fields: Iterable[TargetField]) -> list[FieldDecision]:
    return [classify_field(field) for field in fields if is_v1_target(field)]


def source_evidence_rows(cfg: PlanningConfig) -> list[dict[str, str]]:
    candidates = [
        ("MinStock3 SAP stock", cfg.minstock3_dir / "extract_raw_stock.py", "OITW/OITM SQL for item, warehouse, on hand, minimum, ordered."),
        ("MinStock3 open PO", cfg.minstock3_dir / "extract_raw_stock.py", "OPOR/POR1 SQL for recent open purchase order quantity."),
        ("MinStock3 SPL masters", cfg.minstock3_dir / "map_masters.py", "masters.csv maps linked items to SPLMaster/MasterKey."),
        ("SAP live warehouse lookup", "Service Layer Warehouses endpoint", "WarehouseCode and WarehouseName are fetched live for warehouses present in live stock rows."),
        ("MinStock3 SPI", cfg.minstock3_dir / "enrich_dimensions.py", "SPI_DATA supplies return/repair/list/credit pricing and derived CoCre8 cost."),
        ("MinStock3 usage", cfg.minstock3_dir / "enrich_dimensions.py", "Stock audit DN rows aggregate annual usage."),
        ("Planning V2 live stock extract", "Service Layer SQLQueries on OITW/OITM", "Current stock metrics are fetched live from OITW: OnHand, IsCommited, OnOrder, MinStock, MaxStock, AvgPrice."),
        ("Manual stock audit report", cfg.exco_source_dir / "Stock Audit Report.txt", "Allowed manual SAP report source for usage once mapped to the template."),
        ("Reference target fields", cfg.reference_dir / TARGET_FIELDS_FILE, "Planning V2 target field metadata and CC8 relevance hints."),
        ("Planning V2 sample templates", cfg.samples_dir / "Templates raw.xlsx", "Canonical onboarding template fields used for template-shaped CSV outputs."),
        ("CoCre8 HelpDesk issue tracker", cfg.issue_tracker_csv, "External SharePoint-list export used only for strict ticket/PO evidence; not committed to git."),
        ("HelpDesk Power Automate flow", cfg.reference_dir / "HelpdeskFlow.txt", "Reference flow for ticket lifecycle and improvement recommendations."),
        ("Live SAP item group check", "MinStock3 SAP Service Layer credentials", "Checked Items and ItemGroups while VPN was connected. ItemGroups are broad customer/manufacturer groups such as FUJITSU, ACER, and CHOICE; sampled item master fields ItemType, ItemClass, and MaterialType are generic SAP classifications."),
    ]
    rows = []
    for source_name, path, evidence in candidates:
        exists = path.exists() if isinstance(path, Path) else ""
        rows.append(
            {
                "Source": source_name,
                "Path": str(path),
                "Exists": "Yes" if exists is True else "No" if exists is False else "N/A",
                "Evidence": evidence,
            }
        )
    return rows


def output_object_rows(decisions: Iterable[FieldDecision]) -> list[dict[str, str]]:
    by_object: dict[str, list[FieldDecision]] = {}
    for decision in decisions:
        key = normalize_name(decision.target.object_name or decision.target.context_area)
        by_object.setdefault(key, []).append(decision)

    rows: list[dict[str, str]] = []
    for object_name, fields in CONFIRMED_OUTPUT_OBJECTS.items():
        object_decisions = by_object.get(normalize_name(object_name), [])
        blockers = sorted(
            {
                decision.target.field_name
                for decision in object_decisions
                if decision.source_status in {STATUS_INVESTIGATE_SAP, STATUS_INVESTIGATE_EXTERNAL}
                and priority_rank(decision.target) == 1
            }
        )
        p1_count = sum(1 for decision in object_decisions if priority_rank(decision.target) == 1)
        rows.append(
            {
                "Object": object_name,
                "CSV": f"{object_name}.csv",
                "Readiness": "MVP ready" if p1_count and not blockers else "MVP partial" if p1_count else "Populated",
                "Fields": ", ".join(fields),
                "Blockers": "; ".join(blockers),
            }
        )
    for object_name, fields in PENDING_OUTPUT_OBJECTS.items():
        rows.append(
            {
                "Object": object_name,
                "CSV": f"{object_name}.csv",
                "Readiness": "Pending/header only",
                "Fields": ", ".join(fields),
                "Blockers": "SAP source not proven for v1.",
            }
        )
    return rows


def sap_investigation_rows(decisions: Iterable[FieldDecision]) -> list[dict[str, str]]:
    rows = []
    seen: set[tuple[str, str]] = set()
    candidates = {
        "order_line": ("ServiceCalls / Orders / SQLQueries", "work orders, customer/site, serial, open/resolve timestamps"),
        "service_call_env": ("ServiceCalls / Activity endpoints / SQLQueries", "ETA, recall, resolve, SLA clocks and failure codes"),
        "imo": ("InventoryTransferRequests / StockTransfers / SQLQueries", "from/to warehouse, status, dates, quantities, linked order"),
        "purchase_orders": ("PurchaseOrders / OPOR/POR1 / GoodsReceiptPO", "PO line, vendor, quantity, GRN/receipt status"),
        "repair_orders": ("ServiceCalls / inventory transfers / user tables", "repair order, return authorization, repair status and ETA"),
        "ib_warehouse": ("EquipmentCards / SerialNumbers / BillOfMaterials", "installed base, model, serial, site, BOM quantity"),
        "goods": ("Items / ProductTrees", "model-to-part/BOM relationship"),
        "model": ("Items / ProductTrees", "model ids and descriptions"),
        "person": ("EmployeesInfo / Contacts / SalesPersons", "technician identity and active status"),
        "node": ("Warehouses / Branches / user tables", "warehouse hierarchy and regional node"),
    }
    for decision in decisions:
        object_key = normalize_name(decision.target.object_name)
        if decision.source_status != STATUS_INVESTIGATE_SAP or object_key not in candidates:
            continue
        endpoint, purpose = candidates[object_key]
        key = (object_key, endpoint)
        if key in seen:
            continue
        seen.add(key)
        rows.append(
            {
                "Object": decision.target.object_name,
                "SAP Candidate": endpoint,
                "Purpose": purpose,
                "Status": "Not run",
                "Decision": "Blocked until query/endpoint is proven against CoCre8 SAP.",
            }
        )
    return rows


def unknown_rows(decisions: Iterable[FieldDecision]) -> list[dict[str, str]]:
    rows = []
    for decision in decisions:
        if decision.source_status not in {STATUS_INVESTIGATE_SAP, STATUS_INVESTIGATE_EXTERNAL, STATUS_REVIEW, STATUS_NOT_AVAILABLE}:
            continue
        rows.append(
            {
                "Field": decision.target.field_name,
                "Object": decision.target.object_name,
                "Context": decision.target.context_area,
                "Priority": decision.target.priority,
                "Status": decision.source_status,
                "Question": _investigation_question(decision),
                "Next Action": _next_action(decision),
                "CC8 Comment": decision.target.cc8_comment,
                "Codex": decision.codex_note,
                "Notes": decision.notes,
            }
        )
    return rows


def _investigation_question(decision: FieldDecision) -> str:
    field = decision.target
    if decision.source_status == STATUS_REVIEW:
        return f"Does {field.field_name} mean the same thing for CoCre8 as in the SPL reference workbook?"
    if decision.source_status == STATUS_INVESTIGATE_EXTERNAL:
        return f"Who owns the manual/distribution-portal source for {field.object_name}.{field.field_name}?"
    return f"Which SAP Service Layer endpoint/table supplies {field.object_name}.{field.field_name} for CoCre8?"


def _next_action(decision: FieldDecision) -> str:
    if decision.source_status == STATUS_REVIEW:
        return "Review with business owner before enabling as populated output."
    if decision.source_status == STATUS_INVESTIGATE_EXTERNAL:
        return "Request source file/rule owner, then add parser once received."
    return "Run SAP metadata/query investigation and record result in SAP Investigation Log."


def build_field_map_workbook(cfg: PlanningConfig, output_path: Path | None = None) -> Path:
    fields = parse_target_fields(cfg.reference_dir)
    update_reference_codex_notes(cfg.reference_dir, fields)
    decisions = build_decisions(fields)
    output_path = output_path or Path("docs") / "field-map" / FIELD_MAP_WORKBOOK
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "Target Fields", _target_field_rows(decisions))
    _write_sheet(workbook, "Template Fields", template_field_rows(cfg))
    _write_sheet(workbook, "Source Evidence", source_evidence_rows(cfg))
    _write_sheet(workbook, "Unknowns", unknown_rows(decisions) + template_unknown_rows(cfg))
    _write_sheet(workbook, "Output Objects", output_object_rows(decisions) + template_output_rows(cfg))
    _write_sheet(workbook, "Missing Important Fields", missing_important_field_rows(cfg, decisions))
    _write_sheet(workbook, "SAP Investigation Log", sap_investigation_rows(decisions))
    workbook.save(output_path)
    return output_path


def template_field_rows(cfg: PlanningConfig) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for object_name, fields in load_template_fields(cfg.samples_dir).items():
        populated = POPULATED_TEMPLATE_FIELDS.get(object_name, {})
        for field in fields:
            source = populated.get(field.field_name, "")
            out_of_scope_note = TEMPLATE_FIELDS_NOT_CC8_RELEVANT.get((field.output_object, field.field_name))
            if out_of_scope_note:
                status = STATUS_OUT_OF_SCOPE
                source = ""
                confidence = ""
                notes = out_of_scope_note
            elif source:
                status = STATUS_CONFIRMED
                confidence = "High" if field.field_name not in {"uniqueId", "primaryPartNumber"} else "Medium"
                if field.output_object == "Parts" and field.field_name in {"isPrimary", "primaryPartNumber"}:
                    notes = "Populated from SPI_DATA.csv Main alternative par. Blank when SPI has no main alternative value."
                else:
                    notes = "Populated in template CSV from confirmed live SAP Service Layer or allowed SPI source."
            else:
                status = STATUS_INVESTIGATE_SAP
                confidence = "Low"
                if field.output_object == "Parts" and field.field_name in {"productClass", "productType", "partType"}:
                    status = STATUS_NOT_AVAILABLE
                    notes = "Left blank. Live SAP item master check found broad ItemGroups such as FUJITSU/ACER/CHOICE and generic ItemType/ItemClass/MaterialType values, not Planning V2 class/type semantics."
                elif field.output_object == "Parts":
                    notes = "Left blank. User indicated this field is unlikely to be available from current CoCre8 sources."
                elif field.output_object == "Customers" and field.field_name in CUSTOMER_TEMPLATE_NOTES:
                    status = STATUS_INVESTIGATE_EXTERNAL
                    notes = CUSTOMER_TEMPLATE_NOTES[field.field_name]
                elif field.output_object == "Warehouses" and field.field_name in WAREHOUSE_TEMPLATE_NOTES:
                    status = STATUS_INVESTIGATE_EXTERNAL
                    notes = WAREHOUSE_TEMPLATE_NOTES[field.field_name]
                elif field.output_object == "WarehouseStockOnHand" and field.field_name in STOCK_ON_HAND_TEMPLATE_NOTES:
                    status = STATUS_NOT_AVAILABLE if field.field_name in {"inventoryType", "quantityOutbound"} else STATUS_INVESTIGATE_EXTERNAL
                    notes = STOCK_ON_HAND_TEMPLATE_NOTES[field.field_name]
                elif field.output_object == "PurchaseOrders" and field.field_name == "demandStatus":
                    status = STATUS_REVIEW
                    notes = "Conditional HelpDesk enrichment. Use ReplenishStatus only when PurchaseOrder filename stem matches SAP OPOR.DocNum and part/SPL Master evidence also matches; current export needs reconciliation."
                else:
                    notes = "Left blank in generated template CSV until source and semantics are confirmed."
            rows.append(
                {
                    "Template Object": object_name,
                    "Field": field.field_name,
                    "Field Type": field.field_type,
                    "Template Notes": field.notes,
                    "Source Status": status,
                    "Source": source,
                    "Confidence": confidence,
                    "Implementation Notes": notes,
                }
            )
    return rows


def template_unknown_rows(cfg: PlanningConfig) -> list[dict[str, str]]:
    rows = []
    for row in template_field_rows(cfg):
        if row["Source Status"] in {STATUS_CONFIRMED, STATUS_OUT_OF_SCOPE}:
            continue
        rows.append(
            {
                "Field": row["Field"],
                "Object": row["Template Object"],
                "Context": "Planning V2 template",
                "Status": row["Source Status"],
                "Question": f"What confirmed CoCre8/SAP source should populate {row['Template Object']}.{row['Field']}?",
                "Next Action": "Investigate SAP Service Layer first; only use external/manual source if SAP cannot supply it.",
                "Notes": row["Implementation Notes"],
            }
        )
    return rows


def template_output_rows(cfg: PlanningConfig) -> list[dict[str, str]]:
    rows = []
    by_object: dict[str, list[dict[str, str]]] = {}
    for row in template_field_rows(cfg):
        by_object.setdefault(row["Template Object"], []).append(row)
    for object_name, field_rows in by_object.items():
        relevant = [row for row in field_rows if row["Source Status"] != STATUS_OUT_OF_SCOPE]
        populated = [row for row in relevant if row["Source Status"] == STATUS_CONFIRMED]
        blockers = [row["Field"] for row in relevant if row["Source Status"] != STATUS_CONFIRMED]
        if not relevant:
            readiness = "Out of v1 scope"
        elif not blockers:
            readiness = "Ready"
        elif populated:
            readiness = "Partial"
        else:
            readiness = "Pending/header only"
        rows.append(
            {
                "Object": object_name,
                "CSV": f"{object_name}.csv",
                "Readiness": readiness,
                "Fields": ", ".join(row["Field"] for row in relevant),
                "Blockers": "; ".join(blockers),
            }
        )
    return rows


def _target_field_rows(decisions: Iterable[FieldDecision]) -> list[dict[str, str]]:
    rows = []
    for decision in decisions:
        field = decision.target
        rows.append(
            {
                "Reference Row": str(field.row_number),
                "Field": field.field_name,
                "Object": field.object_name,
                "Context": field.context_area,
                "Used by SPL": field.used_by_spl,
                "CC8 Relevant": field.cc8_relevant,
                "CC8 Supply Hint": field.cc8_can_supply,
                "Priority": field.priority,
                "CC8 Comment": field.cc8_comment,
                "Data Type": field.data_type,
                "Source Status": decision.source_status,
                "Source System": decision.source_system,
                "Source Object": decision.source_object,
                "Source Field": decision.source_field,
                "Transform": decision.transform,
                "Confidence": decision.confidence,
                "Explanation": field.explanation,
                "Codex": decision.codex_note,
                "Notes": decision.notes,
            }
        )
    return rows


def update_reference_codex_notes(reference_dir: Path, fields: list[TargetField]) -> None:
    path = reference_dir / TARGET_FIELDS_FILE
    if not path.exists():
        return
    workbook = load_workbook(path)
    worksheet = workbook["SPL Planning Data Fields"] if "SPL Planning Data Fields" in workbook.sheetnames else workbook.active
    headers = [clean_text(cell.value) for cell in worksheet[1]]
    if "Codex" not in headers:
        codex_col = len(headers) + 1
        worksheet.cell(row=1, column=codex_col, value="Codex")
    else:
        codex_col = headers.index("Codex") + 1
    by_row = {field.row_number: classify_field(field) for field in fields}
    for row_number, decision in by_row.items():
        worksheet.cell(row=row_number, column=codex_col, value=decision.codex_note or decision.notes)
    workbook.save(path)
    workbook.close()


def missing_important_field_rows(cfg: PlanningConfig, decisions: Iterable[FieldDecision]) -> list[dict[str, str]]:
    templates = template_columns(cfg.samples_dir)
    template_fields = {field.lower() for fields in templates.values() for field in fields}
    dictionary_path = cfg.reference_dir / "Planning Data Dictionary (SPL Ref).xlsx"
    dictionary_sheets: list[str] = []
    if dictionary_path.exists():
        workbook = load_workbook(dictionary_path, read_only=True, data_only=True)
        dictionary_sheets = workbook.sheetnames
        workbook.close()

    rows: list[dict[str, str]] = []
    for decision in decisions:
        field = decision.target
        if priority_rank(field) != 1:
            continue
        normalized_field = normalize_name(field.field_name)
        has_template_field = field.field_name.lower() in template_fields or normalized_field in {normalize_name(name) for name in template_fields}
        if has_template_field:
            continue
        rows.append(
            {
                "Potential Gap": f"{field.object_name}.{field.field_name}",
                "Why It Matters": "Priority 1 in edited field list but no direct field-name match was found in the sample templates.",
                "Reference Evidence": f"Context: {field.context_area}; CC8 comment: {field.cc8_comment}",
                "Dictionary Evidence": "Object sheet exists" if normalize_name(field.object_name) in {normalize_name(name) for name in dictionary_sheets} else "No obvious dictionary sheet match",
                "Recommended Action": "Confirm whether this maps to an existing differently named template field or needs a new/onboarding-side evidence field.",
            }
        )

    explicit_gaps = [
        (
            "Part alternatives / SPL Master relationship",
            "Planning operates mainly at SPL Master level and replenishment may use an alternative part.",
            "masters.csv and SPI_DATA.csv prove the relationship, but most transactional templates only have one part field.",
            "Ask Planning V2 whether order/usage/PO line part fields should contain actual part numbers, SPL Master, or both.",
        ),
        (
            "Requested vs dispatched vs replenished part",
            "Ticket lifecycle can involve one requested part, another dispatched alternative, and another replenishment part.",
            "Issue tracker has Part Nr and DispatchPartNo; SAP PO has POR1.ItemCode.",
            "Add explicit evidence fields or confirm importer handles alternatives through Parts.SPLMaster.",
        ),
        (
            "Delivery note number",
            "Delivery note is the strongest bridge between dispatch and actual usage/stock movement.",
            "Stock Audit Report contains DN document; HelpDesk emails often mention attached delivery notes.",
            "Capture delivery-note number in HelpDesk tickets and confirm target template field.",
        ),
        (
            "Fujitsu order confirmation number",
            "Needed to track replenishment after CoCre8 approves PO and SPL orders on Fujitsu portal.",
            "HelpDesk flow already parses OrderConfNumber starting with 800 but templates do not expose it clearly.",
            "Confirm whether it belongs in PurchaseOrders or a separate replenishment/order-lifecycle template.",
        ),
        (
            "Faulty return / collection status",
            "Returned faulty parts are operationally important but unreliable in current process.",
            "Emails contain collection/return instructions; issue tracker has status but not a structured return lifecycle.",
            "Confirm whether this is needed for MVP or only process reporting.",
        ),
        (
            "Stable ticket/conversation identifiers",
            "Strict joins need MSConvoID, Call Number, PO, and line-level part identifiers.",
            "Issue tracker has MSConvoID and Call Number; templates do not obviously carry both.",
            "Confirm whether ServiceOrder/PartsUsage should include request/ticket identifiers.",
        ),
    ]
    for gap, why, evidence, action in explicit_gaps:
        rows.append(
            {
                "Potential Gap": gap,
                "Why It Matters": why,
                "Reference Evidence": evidence,
                "Dictionary Evidence": "Checked against template headers and dictionary sheet names.",
                "Recommended Action": action,
            }
        )
    return rows


def _write_sheet(workbook: Workbook, title: str, rows: list[dict[str, str]]) -> None:
    worksheet = workbook.create_sheet(title)
    headers = list(rows[0].keys()) if rows else ["Note"]
    if not rows:
        rows = [{"Note": "No rows"}]

    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False

    for row in worksheet.iter_rows(min_row=2):
        status = ""
        for cell in row:
            if worksheet.cell(row=1, column=cell.column).value in {"Source Status", "Status", "Readiness"}:
                status = clean_text(cell.value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        fill = _status_fill(status)
        if fill:
            for cell in row:
                cell.fill = fill

    table_ref = f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    table = Table(displayName=re.sub(r"[^A-Za-z0-9_]", "", title)[:25] or "Table", ref=table_ref)
    table.tableStyleInfo = TABLE_STYLE
    worksheet.add_table(table)

    for column_cells in worksheet.columns:
        header = clean_text(column_cells[0].value)
        width = min(max(len(header), *(len(clean_text(cell.value)) for cell in column_cells[1:]), 12), 60)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width + 2


def _status_fill(status: str) -> PatternFill | None:
    if status in {STATUS_CONFIRMED, "Populated"}:
        return CONFIRMED_FILL
    if status in {STATUS_REVIEW, "Partial"}:
        return REVIEW_FILL
    if status in {STATUS_INVESTIGATE_SAP, STATUS_INVESTIGATE_EXTERNAL, STATUS_NOT_AVAILABLE, "Pending/header only"}:
        return UNKNOWN_FILL
    if status == STATUS_OUT_OF_SCOPE:
        return STATIC_FILL
    return None


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the CoCre8 Planning V2 field map workbook.")
    parser.add_argument("--reference-dir", type=Path, default=None)
    parser.add_argument("--cc8-root", type=Path, default=None)
    parser.add_argument("--out", type=Path, default=Path("docs") / "field-map" / FIELD_MAP_WORKBOOK)
    args = parser.parse_args(argv)

    cfg = get_config()
    if args.reference_dir:
        cfg = PlanningConfig(**{**cfg.__dict__, "reference_dir": args.reference_dir})
    if args.cc8_root:
        cfg = PlanningConfig(
            **{
                **cfg.__dict__,
                "cc8_source_root": args.cc8_root,
                "exco_source_dir": args.cc8_root / "Exco" / "source_data",
                "exco_output_dir": args.cc8_root / "Exco" / "output" / "csv",
                "minstock3_dir": args.cc8_root / "MinStock3",
            }
        )
    path = build_field_map_workbook(cfg, args.out)
    print(f"Wrote field map workbook: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
